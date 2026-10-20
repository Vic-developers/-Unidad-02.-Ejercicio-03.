from pathlib import Path
import json
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'Entrega_Ejercicio_03.docx'
metrics = pd.read_csv(ROOT/'reports/text_metrics.csv')
classes = pd.read_csv(ROOT/'reports/class_metrics.csv')
errors = pd.read_csv(ROOT/'reports/error_analysis.csv')
cm = pd.read_csv(ROOT/'reports/confusion_matrix.csv', index_col=0).to_numpy()
selected = metrics.iloc[0]
spam = classes[classes.label=='spam'].iloc[0]
selected_name = {'logistic_regression':'TF-IDF + regresión logística','complement_nb':'TF-IDF + Complement Naive Bayes','dummy':'Baseline Dummy mayoritario'}

D=Document()
sec=D.sections[0]
sec.page_width=Inches(8.5); sec.page_height=Inches(11)
sec.left_margin=Inches(.58); sec.right_margin=Inches(.58)
sec.top_margin=Inches(.48); sec.bottom_margin=Inches(.58)
sec.footer_distance=Inches(.25)
styles=D.styles
normal=styles['Normal']; normal.font.name='Arial'; normal.font.size=Pt(9); normal.font.color.rgb=RGBColor(32,32,32)
normal.paragraph_format.space_after=Pt(4); normal.paragraph_format.line_spacing=1.08
for sname,size in [('Title',24),('Heading 1',21),('Heading 2',12)]:
    s=styles[sname]; s.font.name='Arial'; s.font.size=Pt(size); s.font.bold=True; s.font.color.rgb=RGBColor(25,25,25)
    s.paragraph_format.space_before=Pt(7 if sname!='Title' else 0); s.paragraph_format.space_after=Pt(4)
# Remove Word's built-in Title bottom rule from both style and paragraph formatting.
title_style=styles['Title']
style_ppr=title_style.element.find(qn('w:pPr'))
if style_ppr is not None:
    p_bdr=style_ppr.find(qn('w:pBdr'))
    if p_bdr is not None: style_ppr.remove(p_bdr)

def fmt_para(p, align=WD_ALIGN_PARAGRAPH.JUSTIFY, before=0, after=4, line=1.08):
    p.alignment=align; f=p.paragraph_format; f.space_before=Pt(before); f.space_after=Pt(after); f.line_spacing=line
    return p

def para(text, style=None, align=WD_ALIGN_PARAGRAPH.JUSTIFY, after=4, before=0):
    p=D.add_paragraph(style=style) if style else D.add_paragraph()
    p.add_run(text); return fmt_para(p,align,before,after)

def shade(cell, fill):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); tcPr.append(shd)

def cell_text(cell,text,bold=False,color='202020',size=8.2,align=WD_ALIGN_PARAGRAPH.LEFT):
    cell.text=''; p=cell.paragraphs[0]; p.alignment=align; p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.0
    r=p.add_run(str(text)); r.bold=bold; r.font.name='Arial'; r.font.size=Pt(size); r.font.color.rgb=RGBColor.from_string(color)
    cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER

def set_table_borders(table,color='B8B8B8',sz='5'):
    tblPr=table._tbl.tblPr; borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement(f'w:{edge}'); el.set(qn('w:val'),'single'); el.set(qn('w:sz'),sz); el.set(qn('w:color'),color); borders.append(el)
    tblPr.append(borders)

def fix_table_widths(table,widths):
    table.autofit=False
    grid=table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col=OxmlElement('w:gridCol'); col.set(qn('w:w'),str(round(width*1440))); grid.append(col)
    tblw=table._tbl.tblPr.find(qn('w:tblW'))
    if tblw is None:
        tblw=OxmlElement('w:tblW'); table._tbl.tblPr.append(tblw)
    tblw.set(qn('w:w'),str(round(sum(widths)*1440))); tblw.set(qn('w:type'),'dxa')
    for row in table.rows:
        for idx,cell in enumerate(row.cells): cell.width=Inches(widths[idx])

def set_cell_margins(cell,top=60,start=85,bottom=60,end=85):
    tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        el=OxmlElement('w:'+m); el.set(qn('w:w'),str(v)); el.set(qn('w:type'),'dxa'); mar.append(el)
    tcPr.append(mar)

# Footer
fp=sec.footer.paragraphs[0]; fp.alignment=WD_ALIGN_PARAGRAPH.LEFT
fp.paragraph_format.space_before=Pt(0); fp.paragraph_format.space_after=Pt(0)
r=fp.add_run('INF8239 · Ejercicio 03 · UCI SMS Spam Collection'); r.font.name='Arial'; r.font.size=Pt(7.5); r.font.color.rgb=RGBColor(85,85,85)
fp.add_run('\t')
fp.paragraph_format.tab_stops.add_tab_stop(Inches(7.2))
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); fp._p.append(field)

# Page 1
p=D.add_paragraph(style='Title'); p.add_run('Ejercicio 03'); fmt_para(p,WD_ALIGN_PARAGRAPH.LEFT,0,1,1.0)
direct_ppr=p._p.get_or_add_pPr(); direct_bdr=direct_ppr.find(qn('w:pBdr'))
if direct_bdr is not None: direct_ppr.remove(direct_bdr)
para('Corpus público y clasificador de texto reproducible',after=6)
para('Autor: Victor Elvis Lorenzo     |     Ejecución: 20 octubre 2026',after=2)
para('Repositorio de entrega: https://github.com/Vic-developers/-Unidad-02.-Ejercicio-03..',after=5)
para('Pregunta', 'Heading 2', WD_ALIGN_PARAGRAPH.LEFT, after=3, before=3)
para('¿Cómo se comparan TF-IDF con Complement Naive Bayes y regresión logística para clasificar SMS legítimos (ham) y spam cuando se retiran textos duplicados antes de separar entrenamiento y prueba?',after=5)
# KPI table
kpis=[('5,572','mensajes auditados'),('415','repeticiones retiradas'),('3,867 / 1,290','entrenamiento / prueba'),(f"{selected.f1_macro:.3f}",'F1 macro del mejor')]
t=D.add_table(rows=2,cols=4); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
for i,(value,label) in enumerate(kpis):
    t.columns[i].width=Inches(1.82)
    cell_text(t.cell(0,i),value,True,'202020',14,WD_ALIGN_PARAGRAPH.CENTER); cell_text(t.cell(1,i),label,False,'555555',7.2,WD_ALIGN_PARAGRAPH.CENTER)
    shade(t.cell(0,i),'EAEAEA'); shade(t.cell(1,i),'F3F3F3'); set_cell_margins(t.cell(0,i)); set_cell_margins(t.cell(1,i))
set_table_borders(t)
para('Corpus y decisiones','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3,before=5)
para('SMS Spam Collection, UCI Machine Learning Repository, DOI 10.24432/C5CC84, licencia CC BY 4.0. El ZIP y su huella SHA-256 se incluyen en la entrega. La colección contiene mensajes en inglés con etiquetas ham/spam.')
para('Distribución: 4,825 ham y 747 spam. Sin faltantes ni etiquetas en conflicto. Se detectaron 403 duplicados exactos y 12 coincidencias adicionales tras normalizar mayúsculas/espacios. Se conservó una fila por texto normalizado antes del split; el split fue estratificado 75/25 con semilla 42.')
para('Comparación en el mismo split','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3,before=3)
t=D.add_table(rows=1,cols=3); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
widths=[4.25,1.2,1.2]
for i,w in enumerate(widths): t.columns[i].width=Inches(w)
for i,h in enumerate(['Pipeline','F1 macro','Exactitud']):
    cell_text(t.cell(0,i),h,True,'FFFFFF',8.2); shade(t.cell(0,i),'252525'); set_cell_margins(t.cell(0,i))
for _,row in metrics.iterrows():
    cells=t.add_row().cells
    values=[selected_name.get(row.model,str(row.model)),f'{row.f1_macro:.6f}',f'{row.accuracy:.6f}']
    for i,v in enumerate(values):
        cell_text(cells[i],v,False,'202020',8.1,WD_ALIGN_PARAGRAPH.RIGHT if i else WD_ALIGN_PARAGRAPH.JUSTIFY)
        if len(t.rows)%2==1: shade(cells[i],'EEEEEE')
        set_cell_margins(cells[i])
fix_table_widths(t,widths)
set_table_borders(t)
para(f"Resultado: {selected_name.get(selected.model,selected.model)} fue el mejor en esta partición. En spam, precision={spam['precision']:.3f}, recall={spam['recall']:.3f} y F1={spam['f1']:.3f}; quedan 9 falsos negativos y 14 falsos positivos.",after=2)
p=D.add_paragraph(); fmt_para(p,WD_ALIGN_PARAGRAPH.CENTER,0,0,1.0); p.add_run().add_picture(str(ROOT/'reports/corpus_profile.png'),width=Inches(6.6))

# Page 2
D.add_page_break()
p=D.add_paragraph(style='Heading 1'); p.add_run('Evaluación y lectura de errores'); fmt_para(p,WD_ALIGN_PARAGRAPH.LEFT,0,4,1.0)
para('Resultados observados en una nueva ejecución local; las cifras del informe previo no se reutilizaron como resultados ejecutados.',after=5)
para('Evidencia de la prueba','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3)
para(f"El holdout incluye 1,290 mensajes: 1,129 ham y 161 spam. Matriz (filas = clase real ham/spam; columnas = predicción ham/spam): [[1115, 14], [9, 152]]. El F1 macro de regresión logística es {selected.f1_macro:.6f}, frente a {metrics.loc[metrics.model=='complement_nb','f1_macro'].iloc[0]:.6f} para Complement Naive Bayes y 0.466722 para Dummy.",after=2)
p=D.add_paragraph(); fmt_para(p,WD_ALIGN_PARAGRAPH.CENTER,0,3,1.0); p.add_run().add_picture(str(ROOT/'reports/confusion_matrix.png'),width=Inches(4.3))
para('Revisión cualitativa','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3)
fpn=int(cm[0,1]); fnn=int(cm[1,0]); sample_fp=int(((errors.real=='ham')&(errors.predicted=='spam')).sum()); sample_fn=int(((errors.real=='spam')&(errors.predicted=='ham')).sum())
para(f"El modelo produjo {fpn+fnn} errores en la prueba; se seleccionaron {len(errors)} para análisis: {sample_fp} ham predichos como spam y {sample_fn} spam predichos como ham. Las notas observan mensajes personales muy breves, ambigüedad de expresiones como free/message/offer, promociones con estilo conversacional, y contexto o humor ausente. La tabla está en reports/error_analysis.csv. Las 20 notas son preliminares, requieren revisión del estudiante y no estiman frecuencias poblacionales.")
para('Interpretación y decisión','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3,before=4)
para('La regresión logística superó a Complement Naive Bayes por una diferencia pequeña de 0.0064 puntos de F1 macro. El recall de spam fue 0.944 (9 omisiones). Se conserva como mejor modelo de esta prueba, pero el resultado es un prototipo: un único split del corpus histórico en inglés no permite afirmar rendimiento en mensajes actuales, español dominicano u otros canales.')
para('Reproducibilidad y uso de IA','Heading 2',WD_ALIGN_PARAGRAPH.LEFT,after=3,before=4)
para('El README documenta instalación, dependencias, pruebas y ejecución. La notebook guarda salidas; el código vuelve a calcular los reportes y las figuras. Se ejecutaron 5 pruebas automatizadas, todas aprobadas. Codex se utilizó como apoyo para corregir errores de código, proponer pruebas y revisar la redacción; también ayudó a preparar notas preliminares para interpretar errores. Prompt relevante: integrar LAB04 y LAB05 en una entrega ejecutable; la consigna exige que pueda correrse desde las instrucciones del repositorio. Se verificaron descarga, SHA-256, conteos, partición, métricas, matriz y tests. El estudiante debe revisar las 20 notas y poder explicar el trabajo. El repositorio enlazado contiene los archivos de esta entrega.')
para('Fuentes: UCI Machine Learning Repository, SMS Spam Collection (CC BY 4.0), https://archive.ics.uci.edu/dataset/228/sms+spam+collection. Almeida, Hidalgo y Yamakami (2011), DOI 10.1145/2034691.2034742. Consulte README.md y docs/AI_USAGE.md para instrucciones y declaración ampliada.',after=0)

D.core_properties.title='Ejercicio 03 Corpus público y clasificador de texto reproducible'
D.core_properties.author='Victor Elvis Lorenzo'
D.save(OUT)
print(OUT)
